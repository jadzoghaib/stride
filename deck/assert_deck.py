"""Assert the built deck actually contains the text it is supposed to.

Written after a rebuild silently did not survive the font-embedding step: the
renders were right mid-pipeline and the committed file was a version old. A
check on the final artefact is the only one that counts.

    python assert_deck.py <deck.pptx>
"""

from __future__ import annotations

import pathlib
import shutil
import sys
import tempfile

from pptx import Presentation

MUST_HAVE = [
    "A champion gets given things. Nobody is buying anything.",
    "WHAT THEY DO GET",
    "WHAT EITHER SPORT PAYS HIM",
    "No rate",
]
MUST_NOT_HAVE = [
    "INCOME THE SPORT ITSELF PRODUCED",
    "WHAT EITHER SPORT HAS PAID HIM",
    "Not low. Not inconsistent. Nothing",
]
# A whole shape reading exactly "€0" is the old hero figure. Substring matching
# would also hit "€0.5M" on the pricing slide, which is a real number.
MUST_NOT_BE_A_SHAPE = ["€0"]


def main() -> int:
    src = pathlib.Path(sys.argv[1])
    # Copy first: the live file may be open in PowerPoint or mid-sync.
    with tempfile.TemporaryDirectory() as td:
        tmp = pathlib.Path(td) / "deck.pptx"
        shutil.copy2(src, tmp)
        pres = Presentation(str(tmp))
        texts = [
            sh.text_frame.text
            for slide in pres.slides
            for sh in slide.shapes
            if sh.has_text_frame
        ]
        blob = "\n".join(texts)
        exact = {t.strip() for t in texts}

    bad = [t for t in MUST_HAVE if t not in blob]
    bad += [f"STALE: {t}" for t in MUST_NOT_HAVE if t in blob]
    bad += [f"STALE SHAPE: {t}" for t in MUST_NOT_BE_A_SHAPE if t in exact]
    for b in bad:
        print("  " + b)
    print(f"{src.name}: {len(bad)} problem(s)")
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
