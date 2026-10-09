"""Checks the renders cannot make obvious, run after every build.

Three failures have actually happened in this deck and each one is cheap to
detect mechanically: a shape placed off the canvas (pptxgenjs writes the
coordinate rather than clamping it), a presented slide with no speaker notes,
and a slide carrying no page number. Everything else still needs eyes.

    uv run --with python-pptx python check_deck.py
"""

from __future__ import annotations

import sys

from pptx import Presentation
from pptx.util import Emu

DECK = "Stride_Pitch_Deck.pptx"
W, H = Emu(12192000), Emu(6858000)   # 13.333 x 7.5 inches
PRESENTED = 19                       # slides 1 to 19; the rest are backup


def inches(v: int) -> float:
    return v / 914400


def main() -> int:
    pres = Presentation(DECK)
    problems: list[str] = []
    slides = list(pres.slides)
    print(f"{len(slides)} slides")

    for n, slide in enumerate(slides, 1):
        for sh in slide.shapes:
            if sh.left is None or sh.top is None:
                continue
            right, bottom = sh.left + (sh.width or 0), sh.top + (sh.height or 0)
            if sh.left < 0 or sh.top < 0 or right > W or bottom > H:
                problems.append(
                    f"slide {n}: {sh.shape_type} '{sh.name}' runs off the canvas "
                    f"({inches(sh.left):.2f},{inches(sh.top):.2f}) to "
                    f"({inches(right):.2f},{inches(bottom):.2f})")

        has_notes = (slide.has_notes_slide
                     and slide.notes_slide.notes_text_frame.text.strip())
        if n <= PRESENTED and n != 1 and not has_notes:
            problems.append(f"slide {n}: presented, but carries no speaker notes")

    # Every slide but the covers declares the page-number placeholder on its
    # layout, so the check is that the layout is the one that has it.
    for n, slide in enumerate(slides, 1):
        name = slide.slide_layout.name
        if name == "COVER" and n not in (1, PRESENTED, PRESENTED + 1):
            problems.append(f"slide {n}: on the COVER layout, so it has no page number")

    for p in problems:
        print("  " + p)
    print(f"\n{len(problems)} problem(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
